#!/usr/bin/env python3
"""Compara as sugestões dos testes automáticos (resumo.csv) com as notas registradas na planilha.

Uso:
    python3 analise/comparar_planilha.py resultados-testes/resumo.csv dados/rubrica-tcc.xlsx [saida.csv]

A planilha precisa ter sido recalculada (fórmulas com valores em cache), como a versão publicada.
"""
import csv
import sys

import openpyxl


def main():
    resumo, planilha = sys.argv[1], sys.argv[2]
    saida = sys.argv[3] if len(sys.argv) > 3 else None
    auto = {r["execucao"]: r for r in csv.DictReader(open(resumo, encoding="utf-8-sig"))}
    ws = openpyxl.load_workbook(planilha, data_only=True)["Avaliacao"]
    linhas, iguais, diferentes, manuais = [], 0, 0, 0
    for r in range(3, ws.max_row + 1):
        if ws[f"A{r}"].value is None:
            continue
        nome = f"nivel{ws[f'A{r}'].value}-{ws[f'C{r}'].value}-t{ws[f'E{r}'].value}"
        if nome not in auto:
            continue
        a = auto[nome]
        comp_p, seg_p, rodou_p = ws[f"I{r}"].value, ws[f"L{r}"].value, ws[f"G{r}"].value
        comp_a, seg_a = a["completude_sugerida"], a["seg_soma"]
        if comp_a in ("", "None"):
            situacao, manuais = "manual ou correção", manuais + 1
        else:
            bate = int(float(comp_a)) == int(comp_p)
            if ws[f"A{r}"].value == 4:
                bate = bate and seg_a not in ("", "None") and abs(float(seg_a) - float(seg_p)) < 1e-9
            situacao = "igual" if bate else "DIFERENTE"
            iguais, diferentes = iguais + bate, diferentes + (not bate)
        linhas.append([nome, rodou_p, a["rodou_sem_alteracao"], comp_p, comp_a, seg_p, seg_a, situacao, a["problemas"]])
    print(f"avaliadas automaticamente: {iguais + diferentes} | iguais à planilha: {iguais} | diferentes: {diferentes} | "
          f"manuais ou com correção: {manuais}")
    for l in linhas:
        if l[7] != "igual":
            print("  ", l[0], "|", l[7], "|", l[8][:90])
    if saida:
        with open(saida, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["execucao", "rodou_sem_alteracao_planilha", "rodou_sem_alteracao_teste", "completude_planilha",
                        "completude_teste", "seguranca_planilha", "seguranca_teste", "situacao", "problemas"])
            w.writerows(linhas)


if __name__ == "__main__":
    main()
